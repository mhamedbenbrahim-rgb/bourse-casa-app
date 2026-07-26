# -*- coding: utf-8 -*-
"""
Collecte des tableaux de performances OPCVM publiés par l'ASFIM
(https://asfim.ma/publications/tableaux-des-performances/).

Le site liste les tableaux (quotidiens/hebdomadaires) via l'API JSON
`fundshare.asfim.ma/api/counter/`, chaque tableau étant exportable en
Excel via `fundshare.asfim.ma/api/performances/export/?date=YYYY-MM-DD`.

Ce script est idempotent : chaque exécution ne télécharge que les
tableaux (dates) qui ne figurent pas encore dans la base SQLite locale,
ce qui permet de l'utiliser aussi bien pour la récupération de
l'historique complet (première exécution) que pour le contrôle
quotidien d'un nouveau tableau (exécutions suivantes).

Base : opcvm.db
- table `performances_opcvm` : une ligne par fonds et par date (AN, VL,
  performances glissantes, caractéristiques du fonds).
- table `rapports_traites` : journal des tableaux déjà téléchargés
  (permet de savoir ce qui a déjà été traité et d'éviter les
  retéléchargements).
"""

from __future__ import annotations

import argparse
import io
import sqlite3
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass

import openpyxl
import requests

COUNTER_URL = "https://fundshare.asfim.ma/api/counter/"
EXPORT_URL = "https://fundshare.asfim.ma/api/performances/export/"
DB_PATH = "opcvm.db"
USER_AGENT = (
    "Mozilla/5.0 (compatible; bourse-casa-app/1.0; "
    "+https://github.com/mhamedbenbrahim-rgb/bourse-casa-app)"
)

# En-tête Excel (ligne 2) -> colonne SQLite
COLUMN_MAP = {
    "CODE ISIN": "code_isin",
    "Code Maroclear": "code_maroclear",
    "OPCVM": "opcvm",
    "Société de Gestion": "societe_gestion",
    "Nature juridique": "nature_juridique",
    "Classification": "classification",
    "Sensibilité": "sensibilite",
    "Indice Bentchmark": "indice_benchmark",
    "Périodicité VL": "periodicite_vl",
    "Souscripteurs": "souscripteurs",
    "Affectation des résultats": "affectation_resultats",
    "Commission de souscription": "commission_souscription",
    "Commission de rachat": "commission_rachat",
    "Frais de gestion": "frais_gestion",
    "Dépositaire": "depositaire",
    "Réseau placeur": "reseau_placeur",
    "AN": "an",
    "VL": "vl",
    "YTD": "ytd",
    "1 jour": "perf_1j",
    "1 semaine": "perf_1s",
    "1 mois": "perf_1m",
    "3 mois": "perf_3m",
    "6 mois": "perf_6m",
    "1 an": "perf_1a",
    "2 ans": "perf_2a",
    "3 ans": "perf_3a",
    "5 ans": "perf_5a",
}

NUMERIC_COLUMNS = {
    "commission_souscription", "commission_rachat", "frais_gestion",
    "an", "vl", "ytd", "perf_1j", "perf_1s", "perf_1m", "perf_3m",
    "perf_6m", "perf_1a", "perf_2a", "perf_3a", "perf_5a",
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS performances_opcvm (
    date                    TEXT NOT NULL,
    is_hebdo                INTEGER NOT NULL,
    code_isin               TEXT NOT NULL,
    code_maroclear          TEXT,
    opcvm                   TEXT,
    societe_gestion         TEXT,
    nature_juridique        TEXT,
    classification          TEXT,
    sensibilite             TEXT,
    indice_benchmark        TEXT,
    periodicite_vl          TEXT,
    souscripteurs           TEXT,
    affectation_resultats   TEXT,
    commission_souscription REAL,
    commission_rachat       REAL,
    frais_gestion           REAL,
    depositaire             TEXT,
    reseau_placeur          TEXT,
    an                      REAL,
    vl                      REAL,
    ytd                     REAL,
    perf_1j                 REAL,
    perf_1s                 REAL,
    perf_1m                 REAL,
    perf_3m                 REAL,
    perf_6m                 REAL,
    perf_1a                 REAL,
    perf_2a                 REAL,
    perf_3a                 REAL,
    perf_5a                 REAL,
    PRIMARY KEY (date, code_isin)
);

CREATE INDEX IF NOT EXISTS idx_perf_opcvm_isin ON performances_opcvm(code_isin);
CREATE INDEX IF NOT EXISTS idx_perf_opcvm_date ON performances_opcvm(date);

CREATE TABLE IF NOT EXISTS rapports_traites (
    date         TEXT PRIMARY KEY,
    is_hebdo     INTEGER NOT NULL,
    api_id       INTEGER,
    nb_lignes    INTEGER NOT NULL,
    traite_le    TEXT NOT NULL
);
"""


@dataclass
class Rapport:
    date: str
    is_hebdo: bool
    api_id: int


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update({"User-Agent": USER_AGENT, "Accept": "application/json"})
    return s


def list_available_reports(session: requests.Session) -> list[Rapport]:
    """Récupère la liste complète des tableaux publiés par l'ASFIM."""
    rapports: list[Rapport] = []
    page = 1
    page_size = 1000
    while True:
        resp = session.get(
            COUNTER_URL,
            params={"page": page, "page_size": page_size, "ordering": "date"},
            timeout=30,
        )
        resp.raise_for_status()
        data = resp.json()
        for item in data.get("results", []):
            rapports.append(Rapport(
                date=item["date"], is_hebdo=bool(item["is_hebdo"]),
                api_id=item.get("id"),
            ))
        if not data.get("next"):
            break
        page += 1
    return rapports


def fetch_export(session: requests.Session, date: str, retries: int = 3) -> bytes:
    last_err = None
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(EXPORT_URL, params={"date": date}, timeout=60)
            resp.raise_for_status()
            return resp.content
        except requests.RequestException as exc:
            last_err = exc
            time.sleep(min(2 ** attempt, 10))
    raise RuntimeError(f"Échec du téléchargement pour {date}") from last_err


def parse_export(content: bytes) -> list[dict]:
    """Parse le fichier Excel (ligne 1 = titre, ligne 2 = en-têtes, puis données)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)  # ligne de titre
    header_row = next(rows_iter, None)
    if header_row is None:
        return []
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    columns = [COLUMN_MAP.get(h) for h in headers]

    out = []
    for row in rows_iter:
        if not any(v is not None for v in row):
            continue
        record = {}
        for col_name, value in zip(columns, row):
            if col_name is None:
                continue
            if col_name in NUMERIC_COLUMNS:
                value = value if isinstance(value, (int, float)) else None
            record[col_name] = value
        if record.get("code_isin"):
            out.append(record)
    wb.close()
    return out


def ensure_schema(con: sqlite3.Connection) -> None:
    con.executescript(SCHEMA)


def already_processed_dates(con: sqlite3.Connection) -> set[str]:
    return {row[0] for row in con.execute("SELECT date FROM rapports_traites")}


def upsert_report(con: sqlite3.Connection, rapport: Rapport, records: list[dict]) -> None:
    columns = list(COLUMN_MAP.values())
    placeholders = ", ".join("?" for _ in columns)
    col_list = ", ".join(columns)
    sql = (
        f"INSERT OR REPLACE INTO performances_opcvm "
        f"(date, is_hebdo, {col_list}) VALUES (?, ?, {placeholders})"
    )
    con.executemany(sql, [
        (rapport.date, int(rapport.is_hebdo), *[rec.get(c) for c in columns])
        for rec in records
    ])
    con.execute(
        "INSERT OR REPLACE INTO rapports_traites "
        "(date, is_hebdo, api_id, nb_lignes, traite_le) VALUES (?, ?, ?, ?, datetime('now'))",
        (rapport.date, int(rapport.is_hebdo), rapport.api_id, len(records)),
    )
    con.commit()


def _fetch_and_parse(session: requests.Session, rapport: Rapport) -> tuple[Rapport, list[dict]]:
    content = fetch_export(session, rapport.date)
    return rapport, parse_export(content)


def run(db_path: str = DB_PATH, max_new: int | None = None,
        retry_empty: bool = False, quiet: bool = False, workers: int = 6) -> int:
    """Télécharge et intègre les tableaux non encore traités. Renvoie le nombre traité."""
    con = sqlite3.connect(db_path)
    ensure_schema(con)

    session = make_session()
    reports = list_available_reports(session)
    done = already_processed_dates(con)
    if retry_empty:
        empty_dates = {
            row[0] for row in con.execute(
                "SELECT date FROM rapports_traites WHERE nb_lignes = 0"
            )
        }
        done -= empty_dates

    todo = [r for r in reports if r.date not in done]
    todo.sort(key=lambda r: r.date)
    if max_new is not None:
        todo = todo[:max_new]

    if not quiet:
        print(f"{len(reports)} tableaux publiés, {len(todo)} à traiter.")

    n = 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {
            pool.submit(_fetch_and_parse, session, rapport): rapport
            for rapport in todo
        }
        for future in as_completed(futures):
            rapport = futures[future]
            try:
                rapport, records = future.result()
            except Exception as exc:  # noqa: BLE001 - on journalise et on continue
                print(f"[!] {rapport.date} : {exc}", file=sys.stderr)
                continue
            upsert_report(con, rapport, records)
            n += 1
            if not quiet:
                tag = "hebdo" if rapport.is_hebdo else "quotidien"
                print(f"  {rapport.date} ({tag}) : {len(records)} fonds")

    con.close()
    return n


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", default=DB_PATH, help="Chemin de la base SQLite")
    parser.add_argument("--max-new", type=int, default=None,
                         help="Nombre maximum de nouveaux tableaux à traiter")
    parser.add_argument("--retry-empty", action="store_true",
                         help="Retente les dates déjà journalisées mais sans données")
    parser.add_argument("--workers", type=int, default=6,
                         help="Nombre de téléchargements en parallèle")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()

    n = run(db_path=args.db, max_new=args.max_new,
            retry_empty=args.retry_empty, quiet=args.quiet, workers=args.workers)
    print(f"Terminé : {n} nouveau(x) tableau(x) intégré(s).")


if __name__ == "__main__":
    main()
